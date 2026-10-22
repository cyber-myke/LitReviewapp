"""
make_deck.py
============
Generates an editable PowerPoint (.pptx) tutorial for building and
deploying the Microplastics & Nanoplastics Literature Review Assistant.

Usage
-----
    pip install python-pptx
    python make_deck.py

Output
------
    MNP_Tutorial.pptx  (in the current folder)
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ----------------------------------------------------------------------
# Theme colors
# ----------------------------------------------------------------------
TEAL        = RGBColor(0x0D, 0x73, 0x77)
TEAL_LIGHT  = RGBColor(0x14, 0x91, 0x9B)
DARK        = RGBColor(0x0B, 0x20, 0x27)
PAPER       = RGBColor(0xF7, 0xFA, 0xFA)
INK         = RGBColor(0x1A, 0x2B, 0x30)
MUTED       = RGBColor(0x5B, 0x70, 0x75)
ACCENT      = RGBColor(0xB5, 0x17, 0x9E)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG     = RGBColor(0x0B, 0x20, 0x27)
CODE_FG     = RGBColor(0xD7, 0xF3, 0xF2)
NOTE_BG     = RGBColor(0xFE, 0xF6, 0xEC)
TIP_BG      = RGBColor(0xEE, 0xFA, 0xF9)

# 16:9 canvas
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]  # fully blank layout


# ----------------------------------------------------------------------
# Low-level helpers
# ----------------------------------------------------------------------
def add_slide(bg=PAPER):
    slide = prs.slides.add_slide(BLANK)
    rect = slide.shapes.add_shape(
        1, 0, 0, SLIDE_W, SLIDE_H  # 1 = rectangle
    )
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg
    rect.line.fill.background()
    rect.shadow.inherit = False
    # send background to back
    slide.shapes._spTree.remove(rect._element)
    slide.shapes._spTree.insert(2, rect._element)
    return slide


def add_box(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def style_run(run, text, size, color, bold=False, italic=False,
              font="Segoe UI"):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def add_para(tf, text, size, color, bold=False, italic=False,
             font="Segoe UI", space_after=6, bullet=False, level=0,
             align=PP_ALIGN.LEFT, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.level = level
    p.space_after = Pt(space_after)
    prefix = ""
    if bullet:
        prefix = "▸  " if level == 0 else "–  "
    run = p.add_run()
    style_run(run, prefix + text, size, color, bold, italic, font)
    return p


def colored_panel(slide, left, top, width, height, fill, line=None):
    shp = slide.shapes.add_shape(1, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def accent_bar(slide, left, top, color=TEAL_LIGHT):
    """Thin vertical accent bar for code/note panels."""
    bar = slide.shapes.add_shape(1, left, top, Pt(4), Inches(0.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


# ----------------------------------------------------------------------
# Slide templates
# ----------------------------------------------------------------------
def title_slide(title, subtitle, badge=None, gradient_end=DARK):
    slide = add_slide(bg=TEAL)
    # darker overlay block for depth
    colored_panel(slide, 0, Inches(4.8), SLIDE_W, Inches(2.7), gradient_end)
    _, tf = add_box(slide, Inches(0.9), Inches(1.6),
                    Inches(11.5), Inches(3.0))
    add_para(tf, title, 40, WHITE, bold=True, first=True, space_after=14)
    add_para(tf, subtitle, 19, RGBColor(0xD7, 0xF3, 0xF2))
    if badge:
        _, bf = add_box(slide, Inches(0.9), Inches(5.3),
                        Inches(10), Inches(0.8))
        add_para(bf, badge, 14, WHITE, italic=True, first=True)
    return slide


def section_slide(kicker, title):
    slide = add_slide(bg=DARK)
    _, tf = add_box(slide, Inches(0.9), Inches(2.7),
                    Inches(11.5), Inches(2.5))
    add_para(tf, kicker.upper(), 15, TEAL_LIGHT, bold=True, first=True,
             space_after=12)
    add_para(tf, title, 34, WHITE, bold=True)
    return slide


def content_header(slide, title, subtitle=None):
    _, tf = add_box(slide, Inches(0.8), Inches(0.5),
                    Inches(11.7), Inches(1.3))
    add_para(tf, title, 28, TEAL, bold=True, first=True, space_after=4)
    # underline
    colored_panel(slide, Inches(0.85), Inches(1.32),
                  Inches(11.6), Pt(3), TEAL_LIGHT)
    if subtitle:
        _, sf = add_box(slide, Inches(0.8), Inches(1.45),
                        Inches(11.7), Inches(0.7))
        add_para(sf, subtitle, 15, MUTED, italic=True, first=True)


def bullets_slide(title, subtitle, bullets, note=None, tip=None):
    slide = add_slide()
    content_header(slide, title, subtitle)
    top = Inches(2.3)
    h = Inches(3.4) if (note or tip) else Inches(4.6)
    _, tf = add_box(slide, Inches(0.9), top, Inches(11.5), h)
    first = True
    for b in bullets:
        if isinstance(b, tuple):
            text, level = b
        else:
            text, level = b, 0
        add_para(tf, text, 16 if level == 0 else 14, INK,
                 bullet=True, level=level, first=first, space_after=8)
        first = False
    if note:
        box_callout(slide, note, NOTE_BG, RGBColor(0xE8, 0x5D, 0x04),
                    RGBColor(0x7A, 0x3E, 0x00))
    if tip:
        box_callout(slide, tip, TIP_BG, TEAL_LIGHT, DARK)
    return slide


def box_callout(slide, text, fill, bar_color, text_color,
                top=Inches(5.9)):
    panel = colored_panel(slide, Inches(0.9), top,
                          Inches(11.5), Inches(1.0), fill)
    accent = slide.shapes.add_shape(1, Inches(0.9), top,
                                    Pt(5), Inches(1.0))
    accent.fill.solid()
    accent.fill.fore_color.rgb = bar_color
    accent.line.fill.background()
    accent.shadow.inherit = False
    _, tf = add_box(slide, Inches(1.15), top + Inches(0.1),
                    Inches(11.0), Inches(0.8))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_para(tf, text, 13, text_color, first=True)


def code_slide(title, subtitle, code_lines, note=None, tip=None):
    slide = add_slide()
    content_header(slide, title, subtitle)
    top = Inches(2.3)
    code_h = Inches(3.3) if (note or tip) else Inches(4.6)
    # code panel
    colored_panel(slide, Inches(0.9), top, Inches(11.5), code_h, CODE_BG)
    accent = slide.shapes.add_shape(1, Inches(0.9), top, Pt(5), code_h)
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL_LIGHT
    accent.line.fill.background()
    accent.shadow.inherit = False
    _, tf = add_box(slide, Inches(1.2), top + Inches(0.15),
                    Inches(11.0), code_h - Inches(0.3))
    first = True
    for line in code_lines:
        is_comment = line.strip().startswith("#")
        add_para(tf, line if line else " ", 13,
                 RGBColor(0x6F, 0xA8, 0xA3) if is_comment else CODE_FG,
                 font="Consolas", first=first, space_after=2)
        first = False
    if note:
        box_callout(slide, note, NOTE_BG, RGBColor(0xE8, 0x5D, 0x04),
                    RGBColor(0x7A, 0x3E, 0x00))
    if tip:
        box_callout(slide, tip, TIP_BG, TEAL_LIGHT, DARK)
    return slide


def cards_slide(title, subtitle, cards):
    """cards: list of (heading, body). Lays out in a 2-column grid."""
    slide = add_slide()
    content_header(slide, title, subtitle)
    cols = 2
    cw = Inches(5.6)
    ch = Inches(1.55)
    gx = Inches(0.9)
    gy = Inches(2.4)
    pad_x = Inches(0.3)
    pad_y = Inches(0.25)
    for idx, (head, body) in enumerate(cards):
        r = idx // cols
        c = idx % cols
        left = gx + c * (cw + pad_x)
        top = gy + r * (ch + pad_y)
        colored_panel(slide, left, top, cw, ch, WHITE,
                      line=RGBColor(0xDD, 0xE6, 0xE7))
        _, tf = add_box(slide, left + Inches(0.25), top + Inches(0.15),
                        cw - Inches(0.5), ch - Inches(0.3))
        add_para(tf, head, 15, TEAL, bold=True, first=True, space_after=4)
        add_para(tf, body, 12, INK)
    return slide


def table_slide(title, subtitle, rows, tip=None):
    """rows: list of (filename, description)."""
    slide = add_slide()
    content_header(slide, title, subtitle)
    top = Inches(2.4)
    rh = Inches(0.75)
    for i, (name, desc) in enumerate(rows):
        y = top + i * (rh + Inches(0.12))
        colored_panel(slide, Inches(0.9), y, Inches(3.0), rh, DARK)
        _, nf = add_box(slide, Inches(1.1), y + Inches(0.12),
                        Inches(2.7), rh - Inches(0.24))
        add_para(nf, name, 14, CODE_FG, bold=True, font="Consolas",
                 first=True)
        colored_panel(slide, Inches(4.0), y, Inches(8.4), rh, WHITE,
                      line=RGBColor(0xDD, 0xE6, 0xE7))
        _, df = add_box(slide, Inches(4.2), y + Inches(0.1),
                        Inches(8.0), rh - Inches(0.2))
        df.vertical_anchor = MSO_ANCHOR.MIDDLE
        add_para(df, desc, 13, INK, first=True)
    if tip:
        box_callout(slide, tip, TIP_BG, TEAL_LIGHT, DARK,
                    top=top + len(rows) * (rh + Inches(0.12)) + Inches(0.15))
    return slide


# ======================================================================
# BUILD THE DECK
# ======================================================================

# 1 — Title
title_slide(
    "Building & Deploying the Microplastics & Nanoplastics "
    "Literature Review Assistant",
    "A complete, step-by-step tutorial: from an empty folder to a live "
    "web app — using Python, Streamlit, Biopython, and the PubMed API.",
    "Beginner-friendly  ·  No prior Streamlit experience required",
)

# 2 — What you'll build
cards_slide(
    "What You'll Build",
    "A research dashboard that searches PubMed and turns raw results "
    "into a structured, visual literature review.",
    [
        ("🔍 Search", "Query PubMed by keyword, boolean logic, and year "
                      "range via the NCBI Entrez API."),
        ("📚 Enrich", "Add citation counts and the Relative Citation Ratio "
                      "from the free NIH iCite API."),
        ("🧬 Tag", "Auto-label each paper by polymer type, size class, "
                   "study type, exposure route, and health endpoint."),
        ("⭐ Shortlist", "Mark papers and export only the ones you choose."),
        ("📊 Visualize", "Charts for publications/year, journals, polymer "
                         "mix, citations, plus a keyword network."),
        ("⬇️ Export", "Clean CSV (Excel-safe) and XLSX downloads."),
    ],
)

# 3 — Section: Foundations
section_slide("Part 1", "Foundations — Tools & Setup")

# 4 — Prerequisites
bullets_slide(
    "Prerequisites",
    "Install these once before you start.",
    [
        "Python 3.9+ — download from python.org. On Windows, check "
        "\"Add Python to PATH\" during install.",
        "A code editor — VS Code or PyCharm.",
        "A free NCBI account + API key — get it at "
        "account.ncbi.nlm.nih.gov → Settings → API Key Management.",
        "A GitHub account — required later for live deployment.",
    ],
    tip="An NCBI API key raises your rate limit from 3 to 10 "
        "requests/second. Optional for testing, strongly recommended.",
)

# 5 — Project structure
table_slide(
    "Project Structure",
    "Create a folder — e.g. LitReviewapp — containing three files.",
    [
        ("app.py", "The entire application (all the Python code)."),
        ("requirements.txt", "The libraries the app needs — used locally "
                             "and by the cloud."),
        (".gitignore", "Tells Git which files to keep out of your repo."),
    ],
    tip="This is a single-file application. Everything lives in app.py, "
        "which keeps it simple to understand and deploy.",
)

# 6 — Virtual environment
code_slide(
    "Step 1 — Create a Virtual Environment",
    "Isolates this project's libraries from the rest of your system.",
    [
        "# Create the environment",
        "python -m venv .venv",
        "",
        "# Activate it — Windows (PowerShell)",
        ".venv\\Scripts\\activate",
        "",
        "# Activate it — macOS / Linux",
        "source .venv/bin/activate",
    ],
    tip="When active, your terminal prompt shows (.venv). Everything you "
        "install now stays inside this folder.",
)

# 7 — Dependencies
code_slide(
    "Step 2 — Dependencies",
    "Put these five lines in requirements.txt, then install in one command.",
    [
        "# requirements.txt",
        "streamlit      # the web UI",
        "biopython      # PubMed access",
        "pandas         # data tables",
        "openpyxl       # Excel export",
        "plotly         # interactive charts",
        "",
        "# Install everything:",
        "pip install -r requirements.txt",
    ],
)

# 8 — Section: Code walkthrough
section_slide("Part 2", "Code Walkthrough — How app.py Works")

# 9 — Big picture
bullets_slide(
    "The Big Picture",
    "app.py flows top-to-bottom in six logical blocks.",
    [
        "1. Imports & page config — load libraries, set title and layout.",
        "2. Styling (CSS) — header banner, metric cards, colored pills.",
        "3. Tagging dictionaries & functions — the MNP domain logic.",
        "4. PubMed helpers — build the query, search, parse each field.",
        "5. Analysis helpers — relevance score, citations, chart builders.",
        "6. UI layout — sidebar inputs, search trigger, four result tabs.",
    ],
)

# 10 — Imports
code_slide(
    "Block 1 — Imports & Page Config",
    "Every tool the app uses, and the browser-tab setup.",
    [
        "import io, time, re, json, math",
        "from datetime import datetime",
        "from collections import Counter",
        "from urllib.request import urlopen, Request",
        "",
        "import pandas as pd",
        "import streamlit as st",
        "import plotly.express as px",
        "from Bio import Entrez          # PubMed access",
        "",
        "st.set_page_config(page_title=\"MNP Lit Review\",",
        "                   page_icon=\"🔬\", layout=\"wide\")",
    ],
    tip="Bio.Entrez is the Biopython module that talks to NCBI's "
        "E-utilities — the engine behind the PubMed search.",
)

# 11 — Tagging engine
code_slide(
    "Block 3 — The MNP Tagging Engine",
    "A dictionary of categories → labels → synonyms. This is what makes "
    "it an MNP tool.",
    [
        "DEFAULT_TAXONOMY = {",
        "  \"Polymer Type\": {",
        "     \"Polystyrene (PS)\": [\"polystyrene\", r\"\\bPS\\b\"],",
        "     \"PET\": [\"polyethylene terephthalate\", r\"\\bPET\\b\"],",
        "  },",
        "  \"Health Endpoint\": {",
        "     \"Neurotoxicity\": [\"neurotox\", \"brain\", \"cognitive\"],",
        "  }, ... }",
        "",
        "def tag_text(text, category_dict):",
        "    # scan title+abstract; return matching labels",
        "    ...",
    ],
    tip="Because it's a plain dictionary, users can edit the whole "
        "taxonomy at runtime from the sidebar — no code changes needed.",
)

# 12 — PubMed search
code_slide(
    "Block 4 — Talking to PubMed",
    "Two calls: esearch finds IDs, efetch downloads the records.",
    [
        "def search_pubmed(query, max_results, email, api_key=None):",
        "    Entrez.email = email          # NCBI requires this",
        "    if api_key: Entrez.api_key = api_key",
        "    handle = Entrez.esearch(db=\"pubmed\", term=query,",
        "                 retmax=max_results, sort=\"relevance\")",
        "    record = Entrez.read(handle); handle.close()",
        "    return record[\"IdList\"], int(record[\"Count\"])",
        "",
        "# efetch then pulls full records in batches, sleeping",
        "# between each to respect NCBI rate limits.",
    ],
    note="Without an API key you're capped at 3 requests/second. The code "
         "batches requests and sleeps between them so you're never blocked.",
)

# 13 — Parsing fields
bullets_slide(
    "Block 4 — Parsing Each Paper",
    "Small helper functions dig each field out of PubMed's nested XML.",
    [
        "Bibliographic: title, authors, journal, date, DOI, PMC ID, "
        "affiliation, funding.",
        "Indexed metadata: MeSH terms, author keywords, chemical "
        "substances, publication type, language, country.",
        "Derived: relevance score, abstract word count, open-access flag, "
        "key-findings summary.",
    ],
    tip="Each helper is wrapped in try/except so one malformed record "
        "never crashes the whole run — it's simply skipped.",
)

# 14 — UI layout
cards_slide(
    "Block 6 — The Interface",
    "Sidebar for inputs; four tabs for output.",
    [
        ("Sidebar", "NCBI credentials, keywords, AND/OR logic, year range, "
                    "max results, enrichment toggles, editable taxonomy."),
        ("📄 Results tab", "Filterable paper cards with tag pills, "
                          "shortlist checkboxes, expandable details."),
        ("📊 Dashboard tab", "Five Plotly charts summarizing the result "
                            "set."),
        ("🕸️ Network & ⬇️ Export", "Keyword co-occurrence graph; CSV/Excel "
                                   "download (all or shortlisted)."),
    ],
)

# 15 — Section: Run locally
section_slide("Part 3", "Run It Locally (Testing)")

# 16 — Launch
code_slide(
    "Step 3 — Launch the App",
    "One command starts a local web server.",
    [
        "# Make sure your virtual environment is active, then:",
        "streamlit run app.py",
        "",
        "# Your browser opens at http://localhost:8501",
        "# 1. Enter your NCBI email (and API key) in the sidebar",
        "# 2. Type keywords, e.g. microplastics / nanoplastics",
        "# 3. Set a year range",
        "# 4. Click  🔍 Search PubMed",
    ],
    tip="Edit app.py and save — Streamlit offers \"Rerun\". For code "
        "changes it's safest to stop (Ctrl+C) and relaunch.",
)

# 17 — Troubleshooting
bullets_slide(
    "Common Issues & Fixes",
    "The three most frequent snags when testing locally.",
    [
        "\"No module named …\" → virtual environment isn't active, or deps "
        "not installed. Re-activate and re-run pip install.",
        "Changes don't appear → fully restart: Ctrl+C, then "
        "streamlit run app.py. A browser refresh isn't always enough.",
        "PubMed errors / throttling → confirm your email is entered; add "
        "an NCBI API key to raise the rate limit.",
    ],
    note="Newest papers may show empty MeSH/chemical fields — NCBI hasn't "
         "indexed them yet. This is expected, not a bug.",
)

# 18 — Section: Deploy
section_slide("Part 4", "Deploy Live on Streamlit Cloud")

# 19 — GitHub
code_slide(
    "Step 4 — Put Code on GitHub",
    "Streamlit Cloud deploys from a GitHub repository.",
    [
        "# .gitignore",
        ".venv/",
        "__pycache__/",
        "*.pyc",
        ".streamlit/secrets.toml",
        "",
        "# Push with Git:",
        "git init",
        "git add app.py requirements.txt .gitignore",
        "git commit -m \"Initial commit\"",
        "git branch -M main",
        "git remote add origin <your-repo-url>",
        "git push -u origin main",
    ],
    note="Never commit your NCBI API key. This app takes it as a sidebar "
         "input at runtime, so there's nothing secret to store.",
)

# 20 — Deploy steps
bullets_slide(
    "Step 5 — Deploy",
    "From repo to public URL in four clicks.",
    [
        "Go to share.streamlit.io and sign in with GitHub.",
        "Click New app → Deploy from existing repo.",
        "Select your repository, branch (main), set main file to app.py.",
        "Click Deploy. Streamlit installs from requirements.txt and "
        "launches your app at a public .streamlit.app URL.",
    ],
    tip="Every time you git push a change, the deployed app updates "
        "automatically.",
)

# 21 — Deploy caveat
bullets_slide(
    "One Deployment Caveat",
    "Shared servers share an IP — and PubMed rate-limits by IP.",
    [
        "Streamlit Community Cloud runs your app on a shared server with a "
        "single IP address.",
        "If several users search at once without their own API keys, they "
        "collectively hit NCBI's per-IP limit.",
        "Fix: encourage each user to enter their own NCBI API key — this "
        "ties the rate limit to their key, not the shared server.",
    ],
    tip="For a personal or small-team tool this is rarely a problem; it "
        "only matters at higher concurrent usage.",
)

# 22 — Recap
bullets_slide(
    "Recap — The Full Journey",
    "You now know how to go from zero to a live research tool.",
    [
        "Set up Python, a virtual environment, and dependencies.",
        "Understand the six code blocks: config, styling, tagging, "
        "PubMed, analysis, UI.",
        "Run locally with streamlit run app.py for testing.",
        "Deploy live by pushing to GitHub and connecting Streamlit Cloud.",
    ],
    tip="Next steps: edit the taxonomy for your subfield, add new chart "
        "types, or extend the exporter with a formatted summary report.",
)

# 23 — Closing
title_slide(
    "You're ready to build & ship.",
    "Searchable, taggable, visual, and deployable — a complete MNP "
    "literature review assistant in a single Python file.",
    "A step-by-step tutorial deck",
)

# ----------------------------------------------------------------------
# Save
# ----------------------------------------------------------------------
OUTFILE = "MNP_Tutorial.pptx"
prs.save(OUTFILE)
print(f"Created {OUTFILE} with {len(prs.slides._sldIdLst)} slides.")
print("Open it in PowerPoint, Google Slides, or LibreOffice Impress.")